from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
FIGURES = ROOT / "figures"
SOURCE = ROOT / "报告内容.md"
DOCX = ROOT / "足球联赛管理与在线票务系统课程设计总报告.docx"
TEX = ROOT / "足球联赛管理与在线票务系统课程设计总报告.tex"

INK = "172321"
GREEN = "176548"
PALE = "EAF3EF"
GRAY = "F4F6F5"
LINE = "BAC9C3"
WHITE = "FFFFFF"
AMBER = "E8F0CF"
RED = "B74A43"


def font(size: int, bold: bool = False):
    path = Path("C:/Windows/Fonts/msyh.ttc")
    return ImageFont.truetype(str(path), size, index=0)


def text_size(draw, text, f):
    box = draw.textbbox((0, 0), text, font=f)
    return box[2] - box[0], box[3] - box[1]


def centered(draw, box, label, size=36, color=INK):
    f = font(size)
    x0, y0, x1, y1 = box
    w, h = text_size(draw, label, f)
    draw.text(((x0 + x1 - w) / 2, (y0 + y1 - h) / 2 - 5), label, fill="#" + color, font=f)


def box(draw, rect, label, fill=PALE, outline=LINE, size=34):
    draw.rounded_rectangle(rect, radius=16, fill="#" + fill, outline="#" + outline, width=3)
    centered(draw, rect, label, size)


def arrow(draw, x0, y0, x1, y1, color=GREEN, width=5):
    draw.line((x0, y0, x1, y1), fill="#" + color, width=width)
    if abs(x1 - x0) >= abs(y1 - y0):
        direction = 1 if x1 >= x0 else -1
        draw.polygon([(x1, y1), (x1 - direction * 20, y1 - 11), (x1 - direction * 20, y1 + 11)], fill="#" + color)
    else:
        direction = 1 if y1 >= y0 else -1
        draw.polygon([(x1, y1), (x1 - 11, y1 - direction * 20), (x1 + 11, y1 - direction * 20)], fill="#" + color)


def canvas(width=2000, height=760):
    im = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(im)
    return im, draw


def save_fig(im, name):
    FIGURES.mkdir(parents=True, exist_ok=True)
    im.save(FIGURES / name, dpi=(200, 200), optimize=True)


def make_business_flow():
    im, d = canvas(1600, 1060)
    centered(d, (30, 25, 1570, 95), "一条联赛从报名到赛果的主要流程", 48)
    labels = ["创建赛季", "俱乐部报名", "生成并发布赛程", "用户购票支付", "管理员提交赛果", "积分榜更新"]
    positions = [(75, 205), (615, 205), (1155, 205), (1155, 535), (615, 535), (75, 535)]
    for i, ((x, y), label) in enumerate(zip(positions, labels)):
        box(d, (x, y, x + 370, y + 150), label, PALE if i < 3 else GRAY, size=43 if i != 2 else 38)
    arrow(d, 445, 280, 595, 280)
    arrow(d, 985, 280, 1135, 280)
    arrow(d, 1340, 355, 1340, 515)
    arrow(d, 1155, 610, 1005, 610)
    arrow(d, 615, 610, 465, 610)
    roles = ["赛事管理员", "俱乐部", "系统自动", "普通用户", "赛事管理员", "系统重算"]
    for (x, y), role in zip(positions, roles):
        centered(d, (x, y + 170, x + 370, y + 220), role, 30, GREEN)
    box(d, (70, 865, 760, 995), "报名不足 2 支：取消并通知", "FFF5F2", RED, 34)
    box(d, (840, 865, 1530, 995), "支付超时：释放锁定座位", "FFF8E9", "D8BA74", 34)
    save_fig(im, "business_flow.png")


def make_architecture():
    im, d = canvas(1500, 1290)
    centered(d, (20, 22, 1480, 95), "前后端分离下的 MVC 职责与请求路径", 46)
    tiers = [
        ("视图 View", "Vue 3 · Router · Pinia"),
        ("控制 Controller", "Spring Boot REST Controller"),
        ("业务模型 Model", "Service · 规则 · 事务"),
        ("持久化 Model", "MyBatis Mapper · SQL"),
        ("数据存储", "MySQL · 文件资源"),
    ]
    for i, (title, subtitle) in enumerate(tiers):
        y = 155 + i * 205
        rect = (220, y, 1280, y + 123)
        d.rounded_rectangle(rect, radius=20, fill="#" + (PALE if i in (0, 2, 4) else GRAY), outline="#" + LINE, width=3)
        centered(d, (245, y + 12, 640, y + 100), title, 44)
        centered(d, (640, y + 12, 1250, y + 100), subtitle, 33)
        if i < 4:
            arrow(d, 750, y + 135, 750, y + 188)
    centered(d, (100, 1190, 1400, 1260), "开发：前端 5173 / 后端 8080     部署：JAR 提供页面和 API", 27)
    save_fig(im, "architecture.png")


def make_deployment():
    im, d = canvas(1600, 780)
    centered(d, (30, 25, 1570, 95), "发布包的浏览器、服务与数据部署关系", 46)
    box(d, (80, 235, 400, 420), "浏览器", PALE, LINE, 50)
    box(d, (590, 235, 1010, 420), "Spring Boot JAR", GRAY, LINE, 42)
    box(d, (1200, 235, 1520, 420), "MySQL", PALE, LINE, 48)
    centered(d, (80, 440, 400, 495), "普通用户与管理人员", 29)
    centered(d, (590, 440, 1010, 495), "Vue 页面 · REST API", 29)
    centered(d, (1200, 440, 1520, 495), "业务数据", 29)
    arrow(d, 420, 325, 570, 325)
    arrow(d, 1030, 325, 1180, 325)
    centered(d, (420, 255, 570, 300), "HTTP", 27, GREEN)
    centered(d, (1030, 255, 1180, 300), "数据库连接", 24, GREEN)
    box(d, (625, 565, 975, 685), "上传文件目录", GRAY, LINE, 38)
    arrow(d, 800, 495, 800, 545)
    centered(d, (160, 710, 1440, 765), "各组件可部署在同一台机器；浏览器不直接访问数据库。", 29)
    save_fig(im, "deployment.png")


def make_lifecycle():
    im, d = canvas(1600, 1100)
    centered(d, (20, 20, 1580, 90), "赛季状态随着报名、排赛和比赛推进", 46)
    names = ["草稿", "报名中", "准备中", "进行中", "已结束"]
    positions = [(70, 235), (615, 235), (1160, 235), (1160, 650), (615, 650)]
    for (x, y), name in zip(positions, names):
        box(d, (x, y, x + 360, y + 130), name, PALE, GREEN, 47)
    arrow(d, 430, 300, 595, 300)
    arrow(d, 975, 300, 1140, 300)
    arrow(d, 1340, 390, 1340, 630)
    arrow(d, 1160, 715, 995, 715)
    notes = ["创建后暂不公开", "手动或按时间开启", "排赛发布后等待开赛", "赛季开始后推进", "全部比赛完成"]
    for (x, _), note in zip(positions[:3], notes[:3]):
        centered(d, (x, 155, x + 360, 215), note, 30)
    for (x, _), note in zip(positions[3:], notes[3:]):
        centered(d, (x, 805, x + 360, 865), note, 30)
    arrow(d, 795, 390, 795, 470, RED)
    box(d, (540, 490, 1050, 570), "不足 2 支：取消并通知", "FFF5F2", RED, 33)
    centered(d, (180, 970, 1420, 1055), "已取消的赛季不再排赛，已报名俱乐部可收到取消通知。", 33)
    save_fig(im, "season_lifecycle.png")


def make_data_model():
    im, d = canvas(1600, 1540)
    centered(d, (20, 20, 1580, 95), "核心表的主键、外键与关联方向", 45)

    def entity(x, y, title, table, fields, fill=PALE):
        d.rounded_rectangle((x, y, x + 400, y + 190), radius=17, fill="#" + fill,
                            outline="#" + LINE, width=3)
        d.text((x + 24, y + 14), title, font=font(38), fill="#" + INK)
        d.text((x + 24, y + 63), table, font=font(28), fill="#" + GREEN)
        for i, field in enumerate(fields):
            d.text((x + 24, y + 109 + i * 33), field, font=font(27), fill="#" + INK)

    entity(80, 160, "赛季", "season_info", ["PK  season_id", "状态、开始日期"])
    entity(600, 160, "报名", "club_season_enrollment", ["PK  enrollment_id", "FK  season_id, club_id"])
    entity(1120, 160, "俱乐部", "club_info", ["PK  club_id", "俱乐部名称"])
    entity(80, 500, "轮次", "round_info", ["PK  round_id", "FK  season_id"], GRAY)
    entity(600, 500, "比赛", "match_info", ["PK  match_id", "FK  round_id, season_id"], GRAY)
    entity(1120, 500, "单场座位库存", "match_seat_inventory", ["PK  inventory_id", "FK  match_id"], GRAY)
    entity(600, 840, "订单", "ticket_order", ["PK  order_id", "FK  match_id"])
    entity(1120, 840, "订单明细", "order_item", ["PK  item_id", "FK  order_id, inventory_id"])
    entity(1120, 1180, "电子票", "e_ticket", ["PK  ticket_id", "FK  item_id (唯一)"], GRAY)

    arrow(d, 480, 255, 580, 255)
    arrow(d, 1120, 255, 1020, 255)
    arrow(d, 280, 350, 280, 480)
    arrow(d, 480, 595, 580, 595)
    arrow(d, 1000, 595, 1100, 595)
    arrow(d, 800, 690, 800, 820)
    arrow(d, 1320, 690, 1320, 820)
    arrow(d, 1000, 935, 1100, 935)
    arrow(d, 1320, 1030, 1320, 1160)

    for x0, y0, x1, y1 in [(485, 198, 575, 242), (1025, 198, 1115, 242),
                           (485, 538, 575, 582), (1005, 538, 1095, 582),
                           (1005, 878, 1095, 922)]:
        centered(d, (x0, y0, x1, y1), "1:N", 27, GREEN)
    for x, y, label in [(305, 405, "1:N"), (825, 745, "1:N"),
                        (1345, 745, "1:N"), (1345, 1085, "1:0..1")]:
        d.text((x, y), label, font=font(27), fill="#" + GREEN)

    d.text((82, 900), "PK 主键  |  FK 外键", font=font(30), fill="#" + INK)
    d.text((82, 946), "箭头：主表 → 子表", font=font(30), fill="#" + INK)
    d.text((82, 1010), "报名表限制同一赛季", font=font(28), fill="#" + INK)
    d.text((82, 1048), "同一俱乐部只能报名一次。", font=font(28), fill="#" + INK)
    d.text((82, 1440), "图中只保留关键关系；票区、支付、退款等约束见 schema.sql。", font=font(27), fill="#" + INK)
    save_fig(im, "data_model.png")


def make_ticket_flow():
    im, d = canvas(1600, 1060)
    centered(d, (20, 20, 1580, 90), "购票交易与逐座库存状态", 48)
    names = ["选择票种和人数", "系统挑选连座", "座位锁给待付订单", "模拟支付成功", "生成电子票"]
    positions = [(80, 210), (620, 210), (1160, 210), (1160, 570), (620, 570)]
    for i, ((x, y), name) in enumerate(zip(positions, names)):
        box(d, (x, y, x + 360, y + 145), name, PALE if i % 2 == 0 else GRAY, LINE, 38 if i in (0, 2) else 43)
    arrow(d, 440, 282, 600, 282)
    arrow(d, 980, 282, 1140, 282)
    arrow(d, 1340, 355, 1340, 550)
    arrow(d, 1160, 642, 1000, 642)
    centered(d, (80, 795, 1520, 865), "库存：AVAILABLE  →  LOCKED  →  SOLD", 42)
    box(d, (250, 900, 1350, 1015), "订单超时或取消：LOCKED → AVAILABLE", "FFF5F2", RED, 38)
    save_fig(im, "ticket_flow.png")


def parse_source():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    result = []
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            result.append(("code", "\n".join(code)))
            i += 1
            continue
        if line.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            if match:
                result.append(("figure", match.group(1), match.group(2)))
            i += 1
            continue
        if line.startswith("# "):
            result.append(("h1", line[2:]))
            i += 1
            continue
        if line.startswith("## "):
            result.append(("h2", line[3:]))
            i += 1
            continue
        if line.startswith("### "):
            result.append(("h3", line[4:]))
            i += 1
            continue
        if line.startswith("| "):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("| "):
                cells = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", x) for x in cells):
                    rows.append(cells)
                i += 1
            result.append(("table", rows))
            continue
        if line.startswith("- "):
            items = []
            while i < len(lines) and lines[i].strip().startswith("- "):
                items.append(lines[i].strip()[2:])
                i += 1
            result.append(("bullets", items))
            continue
        para = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "|", "![", "```", "- ")):
            para.append(lines[i].strip())
            i += 1
        result.append(("p", "".join(para)))
    return result


def plain(text):
    return text.replace("**", "").replace("`", "")


def tex_escape(value):
    value = plain(value)
    substitutions = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_\allowbreak{}", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
    return "".join(substitutions.get(c, c) for c in value)


def build_tex(blocks):
    out = [r"\documentclass[UTF8,zihao=-4,a4paper]{ctexart}",
           r"\usepackage[left=2.7cm,right=2.5cm,top=2.6cm,bottom=2.5cm]{geometry}",
           r"\usepackage{fontspec}", r"\setmainfont{Times New Roman}", r"\setCJKmainfont{SimSun}",
           r"\usepackage{graphicx,longtable,array,booktabs,float,fancyhdr,hyperref,caption,needspace}",
           r"\ctexset{section={format=\Large\bfseries\raggedright}}",
           r"\hypersetup{colorlinks=true,linkcolor=black,urlcolor=black}",
           r"\pagestyle{fancy}\fancyhf{}\fancyhead[C]{足球联赛管理与在线票务系统课程设计总报告}\fancyfoot[C]{\thepage}",
           r"\setlength{\headheight}{15pt}\setlength{\parindent}{2em}\setlength{\parskip}{0.3em}\linespread{1.25}",
           r"\renewcommand{\contentsname}{目录}\setcounter{tocdepth}{3}", r"\begin{document}",
           r"\begin{titlepage}\centering\vspace*{3.0cm}",
           r"{\LARGE\bfseries 足球联赛管理与在线票务系统\par}\vspace{0.7cm}",
           r"{\Large 课程设计总报告\par}\vspace{3cm}",
           r"{\large 课程：软件课程设计 I\par}\vspace{1.2cm}",
           r"{\renewcommand{\arraystretch}{1.6}\begin{tabular}{r@{\quad}l}组号：&13\\组员及学号：&肖博文\quad 924106840435\\&徐林时琛\quad 924106840436\\&杨舒念\quad 924106840437\\班级：&9241069501\end{tabular}}",
           r"\vfill\end{titlepage}",
           r"\begingroup\small\linespread{1.0}\selectfont\setlength{\parskip}{0pt}\tableofcontents\endgroup\newpage"]
    title_seen = False
    for block in blocks:
        kind = block[0]
        if kind == "h1":
            if not title_seen:
                title_seen = True
                continue
            label = tex_escape(block[1])
            out.append(r"\section*{" + label + "}" + "\n" + r"\addcontentsline{toc}{section}{" + label + "}")
        elif kind in ("h2", "h3"):
            command = "subsection" if kind == "h2" else "subsubsection"
            label = tex_escape(block[1])
            if kind == "h3" and block[1].startswith(("问题一", "问题二", "问题三")):
                out.append(r"\par\medskip\noindent{\zihao{4}\songti " + label + r"}\par\nobreak\smallskip")
            else:
                out.append("\\" + command + "*{" + label + "}" + "\n" + r"\addcontentsline{toc}{" + command + "}{" + label + "}")
        elif kind == "p":
            out.append(tex_escape(block[1]) + "\n")
        elif kind == "figure":
            name = Path(block[2]).as_posix()
            placement = "[H]" if name in ("figures/data_model.png", "figures/deployment.png") else "[htbp]"
            out.append(r"\begin{figure}" + placement + r"\centering" + "\n" +
                       r"\includegraphics[width=0.94\linewidth]{" + name + "}" + "\n" +
                       r"\caption*{" + tex_escape(block[1]) + "}" + "\n" + r"\end{figure}")
        elif kind == "table":
            rows = block[1]
            count = len(rows[0])
            widths = [0.22, 0.28, 0.40] if count == 3 else [0.27, 0.63]
            spec = "@{}" + "".join(r">{\raggedright\arraybackslash}" + f"p{{{w}\\linewidth}}" for w in widths) + "@{}"
            out.append(r"\Needspace{8\baselineskip}" + "\n" + r"\begin{longtable}{" + spec + r"}\toprule")
            for ridx, row in enumerate(rows):
                cells = [tex_escape(c) for c in row]
                if ridx == 0:
                    cells = [r"\textbf{" + c + "}" for c in cells]
                out.append(" & ".join(cells) + r" \\" + (r"\midrule" if ridx == 0 else ""))
            out.append(r"\bottomrule\end{longtable}")
        elif kind == "bullets":
            out.append(r"\begin{itemize}")
            out += [r"\item " + tex_escape(item) for item in block[1]]
            out.append(r"\end{itemize}")
        elif kind == "code":
            out.append(r"\noindent\begin{minipage}{\linewidth}\small\begin{verbatim}" + "\n" +
                       block[1] + "\n" + r"\end{verbatim}\end{minipage}\par")
    out.append(r"\end{document}")
    TEX.write_text("\n\n".join(out) + "\n", encoding="utf-8")


def set_run_font(run, family="宋体", size=11, bold=False):
    run.font.name = family
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), family)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "目录将在 Word 中更新"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


def style_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.55)
    sec.bottom_margin = Cm(2.45)
    sec.left_margin = Cm(2.65)
    sec.right_margin = Cm(2.45)
    sec.header_distance = Cm(1.2)
    sec.footer_distance = Cm(1.2)
    sec.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Pt(21)
    for style_name, size, before, after in (("Title", 22, 0, 14), ("Heading 1", 17, 15, 9), ("Heading 2", 13, 12, 6), ("Heading 3", 11, 9, 5)):
        style = doc.styles[style_name]
        style.font.name = "黑体"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Pt(0)
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    border = title_ppr.find(qn("w:pBdr"))
    if border is not None:
        title_ppr.remove(border)
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(header.add_run("足球联赛管理与在线票务系统课程设计总报告"), "宋体", 9)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE")


def add_cover(doc):
    doc.add_paragraph("\n\n\n")
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ppr = p._p.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)
    set_run_font(p.add_run("足球联赛管理与在线票务系统"), "黑体", 22, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("课程设计总报告"), "黑体", 18, True)
    doc.add_paragraph("\n\n")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("课程：软件课程设计 I"), "宋体", 12)
    doc.add_paragraph("\n")
    for label, value in (
        ("组号", "13"),
        ("组员及学号", "肖博文  924106840435"),
        ("", "徐林时琛  924106840436"),
        ("", "杨舒念  924106840437"),
        ("班级", "9241069501"),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(11)
        text = f"{label}：{value}" if label else value
        set_run_font(p.add_run(text), "宋体", 12)
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("目录"), "黑体", 18, True)
    toc_entries = [
        ("1 引言", "3"), ("2 需求分析", "3"), ("3 项目设计", "5"),
        ("4 项目开发", "9"), ("5 个人体会", "11"), ("6 大模型使用说明", "12"),
    ]
    doc.add_paragraph("")
    for label, page in toc_entries:
        toc = doc.add_paragraph()
        toc.paragraph_format.first_line_indent = Pt(0)
        toc.paragraph_format.left_indent = Cm(2.5)
        toc.paragraph_format.right_indent = Cm(2.5)
        toc.paragraph_format.space_after = Pt(10)
        toc.paragraph_format.tab_stops.add_tab_stop(Cm(13))
        set_run_font(toc.add_run(label + "\t" + page), "宋体", 12)
    doc.add_page_break()


def shade(cell, value):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), value)
    tc_pr.append(shd)


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    widths = [Cm(3.2), Cm(5.5), Cm(6.8)] if len(rows[0]) == 3 else [Cm(7.5), Cm(8.0)]
    for i, label in enumerate(rows[0]):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        shade(cell, "EAF3EF")
        cell.text = plain(label)
        for p in cell.paragraphs:
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                set_run_font(run, "黑体", 9, True)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = widths[i]
            cells[i].text = plain(value)
            for p in cells[i].paragraphs:
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.2
                for run in p.runs:
                    set_run_font(run, "宋体", 9)
    doc.add_paragraph("")


def build_docx(blocks):
    doc = Document()
    style_document(doc)
    add_cover(doc)
    title_seen = False
    for block in blocks:
        kind = block[0]
        if kind == "h1":
            if not title_seen:
                title_seen = True
                continue
            p = doc.add_paragraph(style="Heading 1")
            set_run_font(p.add_run(plain(block[1])), "黑体", 17, True)
        elif kind in ("h2", "h3"):
            style = "Heading 2" if kind == "h2" else "Heading 3"
            p = doc.add_paragraph(style=style)
            set_run_font(p.add_run(plain(block[1])), "黑体", 13 if kind == "h2" else 11, True)
        elif kind == "p":
            p = doc.add_paragraph()
            value = plain(block[1])
            if value.startswith(("摘要：", "关键词：")):
                p.paragraph_format.first_line_indent = Pt(0)
                prefix, rest = value.split("：", 1)
                set_run_font(p.add_run(prefix + "："), "黑体", 10.5, True)
                set_run_font(p.add_run(rest), "宋体", 10.5)
            else:
                set_run_font(p.add_run(value), "宋体", 10.5)
        elif kind == "figure":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(ROOT / block[2]), width=Cm(15.5))
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.first_line_indent = Pt(0)
            caption.paragraph_format.space_after = Pt(10)
            set_run_font(caption.add_run(block[1]), "宋体", 9)
        elif kind == "table":
            add_table(doc, block[1])
        elif kind == "bullets":
            for item in block[1]:
                p = doc.add_paragraph(style="List Bullet")
                set_run_font(p.add_run(plain(item)), "宋体", 10.5)
        elif kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_after = Pt(8)
            set_run_font(p.add_run(block[1]), "Consolas", 9)
    doc.save(DOCX)


def main():
    for builder in (make_business_flow, make_architecture, make_deployment, make_lifecycle, make_data_model, make_ticket_flow):
        builder()
    blocks = parse_source()
    build_tex(blocks)
    print(f"Written: {TEX}")
    print(f"Blocks: {len(blocks)}")


if __name__ == "__main__":
    main()
